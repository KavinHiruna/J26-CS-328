"""
expert_validation.py - EVIDENCE that the ranking agrees with people.

Two commands (from the dev folder):
  python -m backend.component3_soc.analysis.expert_validation sheet      make a BLIND rating sheet for a new rater
  python -m backend.component3_soc.analysis.expert_validation            evaluate the sheets in data/expert/

The blind sheet shows only raw facts (attack type, confidence, target role, recent repeats), not our score, level or
severity, so the rater is not influenced. Ratings: Critical = 4 ... Low = 1.
Evaluation compares three scoring methods with the raters' average rating (Spearman, top-20 overlap), reports how much
the raters agree with EACH OTHER (the ceiling for any method) and fits a small cross-validated decision tree.
"""
import glob
import sys
from itertools import combinations
from pathlib import Path

import numpy as np
import pandas as pd
from scipy.stats import spearmanr
from sklearn.metrics import cohen_kappa_score
from sklearn.model_selection import KFold, cross_val_predict
from sklearn.tree import DecisionTreeRegressor, export_text

from .. import pipeline
from ..config import EXPERT_DIR, EXPERT_KEY_FILE, OUTPUT_DIR, REPORT_DIR, WEIGHTS

LEVEL = {"Critical": 4, "High": 3, "Medium": 2, "Low": 1}
KEYS = ["severity", "confidence", "asset", "frequency"]


def target_role(asset_name):
    if asset_name.startswith("Control endpoint"):
        return "Control endpoint (Modbus, MQTT and web services)"
    if asset_name.startswith("Sensor"):
        return "Sensor or monitoring client"
    return "Device not in the inventory"


def build_blank_sheet():
    """Blank rating sheet for the 100 alerts in data/expert/expert_sample_key.csv (same alerts for every rater)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    _, queue = pipeline.run(save=False)
    key = pd.read_csv(EXPERT_KEY_FILE)
    sample = key.merge(queue, on="flow_id")
    wb = Workbook()
    ws0 = wb.active
    ws0.title = "Instructions"
    for i, (text, bold) in enumerate([
            ("Which alert should a SOC analyst handle first?  Rating sheet", True), ("", False),
            ("Context: an industrial plant. An intrusion detector raised the 100 alerts on the next sheet.", False),
            ("For EACH alert choose how urgently a SOC analyst should handle it (dropdown in the Rating column):", False),
            ("    Critical = handle immediately   High = handle soon   Medium = handle when capacity allows   Low = can wait", False),
            ("Use your own judgement; work alone; about 30 minutes. Do NOT use any AI tool.", False),
            ("Put your years of security experience below, and send the file back changed only in Rating/Comment.", False),
            ("Years of security experience:", True)], 1):
        ws0.cell(i, 1, text).font = Font(bold=bold)
    ws0.column_dimensions["A"].width = 120
    ws = wb.create_sheet("Alerts")
    heads = ["No.", "Attack type", "Detector confidence", "Target device", "Repeats in last minute", "Rating (choose)", "Comment (optional)"]
    for j, h in enumerate(heads, 1):
        ws.cell(1, j, h).font = Font(bold=True)
        ws.cell(1, j).fill = PatternFill("solid", fgColor="DDE6F0")
        ws.cell(1, j).alignment = Alignment(wrap_text=True)
    for i, r in enumerate(sample.sort_values("sheet_no").itertuples(), 2):
        repeats = int(r.window_count) if str(r.frequency_basis).startswith("time") and pd.notna(r.window_count) else "n/a"
        for j, v in enumerate([int(r.sheet_no), r.predicted_type.replace("_", " "), float(r.detection_confidence),
                               target_role(r.asset_name), repeats], 1):
            ws.cell(i, j, v)
        ws.cell(i, 3).number_format = "0%"
    validation = DataValidation(type="list", formula1='"Critical,High,Medium,Low"', allow_blank=False)
    ws.add_data_validation(validation)
    validation.add(f"F2:F{len(sample) + 1}")
    for col, width in zip("ABCDEFG", [6, 22, 12, 46, 14, 16, 40]):
        ws.column_dimensions[col].width = width
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    path = OUTPUT_DIR / "expert_rating_sheet_blank.xlsx"
    wb.save(path)
    print("Saved", path, "- send one copy to each rater")


def evaluate():
    files = sorted(glob.glob(str(EXPERT_DIR / "*.xlsx")))
    if not files:
        raise SystemExit("No rating sheets in data/expert/.")
    key = pd.read_csv(EXPERT_KEY_FILE)
    ratings = {}
    for path in files:
        sheet = pd.read_excel(path, sheet_name="Alerts")
        column = [c for c in sheet.columns if str(c).startswith("Rating")][0]
        ratings[Path(path).stem] = sheet.set_index("No.")[column].map(LEVEL)
    R = pd.DataFrame(ratings).loc[key.sheet_no]
    R.index = key.flow_id
    _, queue = pipeline.run(save=False)
    S = queue.set_index("flow_id").loc[key.flow_id]
    f4 = pd.DataFrame({"severity": (S.attack_severity - 1) / 4, "confidence": S.detection_confidence,
                       "asset": (S.asset_criticality - 1) / 4, "frequency": (S.attack_frequency - 1) / 4})
    methods = {"Risk score (AHP weights)": S.risk_score, "Severity only": S.attack_severity,
               "Equal weights (25% each)": 100 * f4.mean(axis=1)}
    mean_rating = R.mean(axis=1)
    lines = [f"RATERS: {len(files)}   alerts rated: {int(R.notna().all(axis=1).sum())}/{len(R)}"]
    lines += [f"  {name}" for name in R.columns]
    rows = []
    top = lambda v, k=20: set(v.nlargest(k).index)
    for name, values in methods.items():
        rho, p = spearmanr(values, mean_rating)
        per = [spearmanr(values, R[c])[0] for c in R.columns]
        rows.append({"method": name, "spearman_with_raters_mean": round(rho, 3),
                     "per_rater_min": round(min(per), 3), "per_rater_max": round(max(per), 3),
                     "top20_overlap": round(len(top(values) & top(mean_rating)) / 20, 2)})
    lines += ["", pd.DataFrame(rows).to_string(index=False)]
    between = []
    if len(files) >= 2:
        for a, b in combinations(R.columns, 2):
            rho = spearmanr(R[a], R[b])[0]
            between.append(rho)
            kappa = cohen_kappa_score(R[a], R[b], weights="quadratic")
            lines.append(f"\nAgreement between {a} and {b}: Spearman {rho:.2f}, weighted kappa {kappa:.2f}  (the ceiling for any method)")
    tree = DecisionTreeRegressor(max_depth=3, min_samples_leaf=5, random_state=42)
    predicted = cross_val_predict(tree, f4.to_numpy(), mean_rating.to_numpy(), cv=KFold(5, shuffle=True, random_state=42))
    tree.fit(f4.to_numpy(), mean_rating.to_numpy())
    lines += [f"\nSmall decision tree learned from the raters' average (depth 3, 5-fold cross-validated): "
              f"Spearman {spearmanr(predicted, mean_rating)[0]:.3f}", export_text(tree, feature_names=KEYS),
              "Reading: a higher Spearman and top-20 overlap means closer to the raters. If the AHP score is not clearly above "
              "equal weights, say so: the weights then add transparency, not accuracy."]
    text = "\n".join(lines)
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    (REPORT_DIR / "expert_agreement.txt").write_text(text, encoding="utf-8")
    print(text)
    return {"ok": True, "methods": pd.DataFrame(rows).set_index("method"), "between_raters": between, "raters": len(files)}


def main():
    if len(sys.argv) > 1 and sys.argv[1] == "sheet":
        build_blank_sheet()
    else:
        return evaluate()


if __name__ == "__main__":
    main()
